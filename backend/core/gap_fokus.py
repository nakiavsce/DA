"""gap_fokus — DATA_YANG_PERLU_DIISI_DA_FOKUS.xlsx: berkas isian yang HANYA memuat yang masih perlu diisi klien
(BOM_AKSESORIS · VARIAN_BARU · MATERIAL). Semua kolom lain sudah terisi dari sistem; klien hanya mengisi sel KUNING.
Sel BIRU = sudah diisi otomatis oleh sistem (mis. varian dari nama bahan pembeda) — cukup diperiksa.
Berkas ini bisa diunggah balik apa adanya ke layar yang sama (kolom-kolom importir tidak diubah posisinya).
"""
from __future__ import annotations

import difflib
import io
import re

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from core import gap_sisa
from core.bom_fill import CATEGORIES, _SPLIT_RE, _variants_label, clean, load_model_variants, match_tokens, target_variants
from core.master_fill import BOM_COLS, MAT_COLS, _head

YELLOW = PatternFill("solid", fgColor="FFF2CC")   # wajib diisi
BLUE = PatternFill("solid", fgColor="DDEBF7")     # diisi otomatis — periksa
GREY = PatternFill("solid", fgColor="EDEDED")     # tidak perlu diubah
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")

BOM_FOKUS_COLS = BOM_COLS + ["yang_perlu_diisi"]
VARIAN_FOKUS_COLS = gap_sisa.VARIAN_BARU_COLS
MAT_FOKUS_COLS = MAT_COLS
_SIZE_RE = re.compile(r"\b(?:size|uk|ukuran)\s*[:.]?\s*(xxl|xl|l|m|s|allsize|all size|std|jmb)\b", re.I)

PETUNJUK = [
    ("DATA YANG MASIH PERLU DIISI — FOKUS (BOM · VARIAN BARU · MATERIAL)", True),
    ("Berkas ini hanya berisi yang masih kurang. Semua baris sudah terisi dari sistem — Anda cukup mengisi sel KUNING.", False),
    ("", False),
    ("WARNA SEL", True),
    ("  KUNING  = wajib Anda isi.", False),
    ("  BIRU    = sudah diisi otomatis oleh sistem (mis. varian ditebak dari nama bahan 'Kancing … warna Mahogany'). Periksa; ubah bila salah.", False),
    ("  ABU-ABU = tidak perlu diubah (sudah benar / diselesaikan di sheet lain).", False),
    ("", False),
    ("URUTAN KERJA (3 langkah)", True),
    ("  1. Sheet VARIAN_BARU — model yang belum punya SKU: isi 'ukuran' (pilih dari ukuran_tersedia). kode_warna sudah disarankan (biru).", False),
    ("     SKU MODEL-WARNA-UKURAN + barang jadinya dibuat otomatis saat diunggah; harga_jual boleh dikosongkan.", False),
    ("  2. Sheet BOM_AKSESORIS — kelompok bahan yang MASIH butuh isian Anda (ada sel kuning). Kolom paling kanan 'yang_perlu_diisi' menyebut persis apa yang kurang di baris itu.", False),
    ("     Baris ber-kode_model = awal kelompok; baris di bawahnya (kode_model kosong) = bahan lain kelompok yang sama. Jangan hapus baris yang sudah benar.", False),
    ("     Kolom 'varian' = warna/ukuran pemakai bahan itu, dipisah koma, pilih dari 'varian_tersedia'. Kosong = semua varian model.", False),
    ("     Sheet BOM_OTOMATIS — kelompok yang sudah diisi otomatis oleh sistem (varian biru) atau selesai lewat sheet lain (VARIAN_BARU/MATERIAL).", False),
    ("     Tidak ada sel kuning di sana: cukup periksa. Sheet ini IKUT DITERAPKAN saat diunggah balik — biarkan apa adanya, jangan dihapus.", False),
    ("  3. Sheet MATERIAL — bahan yang harganya masih 0 atau isi kemasannya belum diketahui. Isi 'harga_per_satuan_beli' (harga 1 roll / 1 m / 1 pack)", False),
    ("     dan, bila diminta, 'isi_per_satuan_beli' = berapa pcs dalam 1 roll/pack. satuan_beli sudah diisi sama dengan satuan dasar.", False),
    ("", False),
    ("UNGGAH BALIK berkas ini di Portal Keuangan → Master Akuntansi → Impor Harga · Rekening · BOM → 'Terapkan semua sheet'.", True),
    ("  Baris yang sel kuningnya dibiarkan kosong TIDAK diubah — aman diunggah sebagian, lalu unduh lagi berkas FOKUS untuk sisanya.", False),
    ("  Kode bahan lihat sheet REF_AKSESORIS (hanya referensi, tidak perlu diisi).", False),
]


def _norm(s) -> str:
    return re.sub(r"[^a-z0-9]", "", clean(s).lower())


def _widths(ws, widths):
    for i, w in enumerate(widths):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w


def _fill(ws, row_idx: int, cols: list[str], names: list[str], fill: PatternFill) -> None:
    for n in names:
        ws.cell(row=row_idx, column=cols.index(n) + 1).fill = fill


# ═══════════════════════════════════════════════════════════════════════════
# SARAN VARIAN — dari nama bahan PEMBEDA antar-kelompok satu model (bukan tebakan importir)
# ═══════════════════════════════════════════════════════════════════════════
def suggest_varian(distinct_names: list[str], variants: list[dict]) -> tuple[str, str]:
    """(saran, dasar). Saran hanya bila TEPAT satu warna (dan ≤1 ukuran) ditemukan pada nama bahan pembeda."""
    text = " ".join(distinct_names)
    ntext = _norm(text)
    words = set(re.findall(r"[a-z0-9]+", text.lower()))
    colors: dict[str, str] = {}
    for v in variants:
        name = clean(v.get("color_name") or "")
        code = (v.get("color_code") or "").upper()
        if len(_norm(name)) >= 3 and _norm(name) in ntext:
            colors[name] = name
        elif len(code) >= 3 and code.lower() in words:
            colors[name or code] = code
        elif len(_norm(name)) >= 4 and " " not in name.strip() and difflib.get_close_matches(_norm(name), [w for w in words if len(w) >= 4], n=1, cutoff=0.85):
            colors[name] = difflib.get_close_matches(_norm(name), [w for w in words if len(w) >= 4], n=1, cutoff=0.85)[0]  # typo: Maron → MAROON
    sizes_avail = {(v.get("size_code") or "").upper() for v in variants}
    sizes = {m.group(1).upper().replace(" ", "") for m in _SIZE_RE.finditer(text)} & sizes_avail
    if len(colors) != 1 or len(sizes) > 1:
        return "", ""
    color = next(iter(colors))
    parts = [color] + sorted(sizes)
    hit = next((n for n in distinct_names if _norm(colors[color]) in _norm(n)), distinct_names[0] if distinct_names else "")
    return ", ".join(parts), hit


def _groups_in_file(rows: list[tuple]) -> dict[int, dict]:
    """head_row → {model_code, codes:set, names:{code:name}} dari baris BOM_AKSESORIS berkas."""
    heads = gap_sisa.group_header_rows(rows)
    out: dict[int, dict] = {}
    for r_idx, h in heads.items():
        src = rows[r_idx - 2]
        g = out.setdefault(h, {"model_code": clean(rows[h - 2][0]).upper(), "codes": set(), "names": {}})
        code = clean(src[2]).upper()
        if code:
            g["codes"].add(code)
            g["names"][code] = clean(src[3])
    return out


async def _variants_by_code(db, model_codes: set[str]) -> dict[str, list[dict]]:
    models = await db.rahaza_models.find({"code": {"$in": list(model_codes)}}, {"_id": 0, "id": 1, "code": 1}).to_list(5000)
    vmap = await load_model_variants(db, [m["id"] for m in models])
    return {m["code"]: vmap.get(m["id"]) or [] for m in models}


async def bom_fokus_rows(db, parsed: dict, rows: list[tuple]) -> tuple[list[list], list[dict], dict]:
    """→ (baris sheet, gaya per baris [{row_i, fills:[(kolom, fill)]}], stats)."""
    heads = gap_sisa.group_header_rows(rows)
    groups = _groups_in_file(rows)
    by_head: dict[int, list[dict]] = {}
    for i in parsed.get("bom_issues") or []:
        if i["kategori"] == "model_dihentikan":
            continue
        by_head.setdefault(heads.get(i["row"], i["row"]), []).append(i)
    vars_by_code = await _variants_by_code(db, {g["model_code"] for g in groups.values()})
    # bahan yang ada di SEMUA kelompok satu model = bukan pembeda
    per_model: dict[str, list[set]] = {}
    for g in groups.values():
        per_model.setdefault(g["model_code"], []).append(g["codes"])
    common = {mc: set.intersection(*sets) if len(sets) > 1 else set() for mc, sets in per_model.items()}

    out: list[list] = []
    styles: list[dict] = []
    stats = {"kelompok": 0, "varian_disarankan": 0, "varian_manual": 0, "baris": 0}
    for h in sorted(by_head):
        members = [r for r in sorted(heads) if heads[r] == h] or [h]
        g = groups.get(h) or {"model_code": "", "codes": set(), "names": {}}
        variants = vars_by_code.get(g["model_code"]) or []
        vlabel = _variants_label(variants) if variants else ""
        group_issues = [i for i in by_head[h] if i["row"] == h and i["kategori"] in ("kelompok_tanpa_varian", "warna_tak_dikenal", "model_tanpa_sku", "model_tak_dikenal")]
        stats["kelompok"] += 1
        saran, dasar = "", ""
        if any(i["kategori"] == "kelompok_tanpa_varian" for i in group_issues) and variants:
            distinct = [g["names"][c] for c in sorted(g["codes"] - common.get(g["model_code"], set()))]
            saran, dasar = suggest_varian(distinct, variants)
            if saran:
                colors, sizes, _un = match_tokens([t for t in _SPLIT_RE.split(saran) if t.strip()], variants)
                if not target_variants({"targets": {"colors": sorted(colors), "sizes": sorted(sizes)}}, variants):
                    saran = ""
        for r in members:
            src = rows[r - 2] if 0 <= r - 2 < len(rows) else (None,) * 9
            row_issues = [i for i in by_head[h] if i["row"] == r and i not in group_issues]
            varian_val = src[7]
            fills: list[tuple[str, PatternFill]] = []
            todo: list[str] = []
            if r == h:
                for i in group_issues:
                    k = i["kategori"]
                    if k == "kelompok_tanpa_varian":
                        if saran:
                            varian_val = saran
                            fills.append(("varian", BLUE))
                            todo.append(f"varian diisi otomatis dari bahan pembeda '{dasar}' → periksa; ubah bila salah")
                            stats["varian_disarankan"] += 1
                        else:
                            fills.append(("varian", YELLOW))
                            todo.append("isi kolom varian: pilih warna/ukuran dari varian_tersedia (pisah koma)")
                            stats["varian_manual"] += 1
                    elif k == "warna_tak_dikenal":
                        fills.append(("varian", YELLOW))
                        todo.append(f"ganti '{i.get('token')}' di kolom varian dengan salah satu dari varian_tersedia")
                    elif k == "model_tanpa_sku":
                        fills.append(("varian", GREY))
                        todo.append(f"tidak ada yang diisi di sini — isi 'ukuran' model {g['model_code']} di sheet VARIAN_BARU; kelompok ini otomatis ikut")
                    elif k == "model_tak_dikenal":
                        fills.append(("kode_model", YELLOW))
                        todo.append("kode model tidak ada di master — perbaiki kode_model")
            for i in row_issues:
                k, det = i["kategori"], i.get("detail", "")
                if k == "qty_kosong":
                    fills.append(("qty_per_pcs", YELLOW))
                    todo.append("isi qty_per_pcs (contoh: 1 pcs, 60 cm, 0,5)")
                elif k == "satuan_tak_valid" and "isi_per_satuan_beli" in det:
                    fills.append(("qty_per_pcs", GREY))
                    todo.append(f"baris ini tidak diubah — isi 'isi_per_satuan_beli' (pcs per {i.get('unit_raw') or 'kemasan'}) untuk {i.get('code')} di sheet MATERIAL")
                elif k == "satuan_tak_valid":
                    m = re.search(r"satuan dasar '([^']+)'", det)
                    fills.append(("qty_per_pcs", YELLOW))
                    todo.append(f"tulis qty dengan satuan yang cocok dengan satuan dasar '{m.group(1) if m else '?'}' (contoh: 60 cm, 0,5 m)")
                elif k == "kode_tak_dikenal":
                    fills.append(("kode_material", YELLOW))
                    mm = re.search(r"mirip: ([A-Z0-9-]+)", det)
                    todo.append("ganti kode_material dengan kode di REF_AKSESORIS" + (f" (mungkin {mm.group(1)})" if mm else ""))
                elif k == "baris_tanpa_model":
                    fills.append(("kode_model", YELLOW))
                    todo.append("tulis kode_model di baris pertama kelompok")
                else:
                    todo.append(f"{CATEGORIES.get(k, k)}: {det}")
            out.append([src[0], src[1], src[2], src[3], src[4], src[5], src[6] if r != h else (src[6] or ""), varian_val,
                        vlabel if r == h else "", "; ".join(dict.fromkeys(todo)) or ("" if r != h else "sudah benar — biarkan")])
            styles.append({"fills": fills, "group": h})
            stats["baris"] += 1
    return out, styles, stats


def split_otomatis(rows: list[list], styles: list[dict]) -> tuple[list[tuple[list, dict]], list[tuple[list, dict]]]:
    """(manual, otomatis). Kelompok tanpa satu pun sel KUNING = tidak ada yang perlu diisi klien → sheet BOM_OTOMATIS."""
    manual_groups = {st["group"] for st in styles if any(f is YELLOW for _c, f in st["fills"])}
    pairs = list(zip(rows, styles))
    return [p for p in pairs if p[1]["group"] in manual_groups], [p for p in pairs if p[1]["group"] not in manual_groups]


async def build_fokus_workbook(db, parsed: dict | None, data: bytes | None) -> tuple[bytes, dict]:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PETUNJUK"
    for text, bold in PETUNJUK:
        ws.append([text])
        if bold:
            ws.cell(row=ws.max_row, column=1).font = BOLD
    ws.column_dimensions["A"].width = 150
    stats: dict = {}

    # ── VARIAN_BARU ──
    ws = wb.create_sheet("VARIAN_BARU")
    _head(ws, VARIAN_FOKUS_COLS)
    if parsed:
        vb = await gap_sisa.varian_baru_rows(db, parsed)
    else:  # tanpa berkas: model aktif yang belum punya satu pun varian (bukan yang dihentikan)
        no_sku = []
        for m in await db.rahaza_models.find({"active": {"$ne": False}}, {"_id": 0, "id": 1, "code": 1, "name": 1}).sort("code", 1).to_list(5000):
            if await db.rahaza_model_variants.count_documents({"model_id": m["id"]}) == 0:
                no_sku.append({"kategori": "model_tanpa_sku", "row": 0, "detail": "", "model_code": m["code"], "model_name": m["name"], "varian_raw": ""})
        vb = await gap_sisa.varian_baru_rows(db, {"bom_issues": no_sku})
    for r in vb:
        ws.append(r)
        i = ws.max_row
        _fill(ws, i, VARIAN_FOKUS_COLS, ["kode_warna"], BLUE if r[3] else YELLOW)
        _fill(ws, i, VARIAN_FOKUS_COLS, ["ukuran"], YELLOW)
        ws.cell(row=i, column=VARIAN_FOKUS_COLS.index("keterangan") + 1).value = \
            ("isi 'ukuran' (pilih dari ukuran_tersedia); kode_warna sudah disarankan — periksa" if r[3]
             else "isi 'kode_warna' (pilih dari kode_warna_tersedia) dan 'ukuran'")
    _widths(ws, (12, 22, 16, 12, 12, 14, 62, 90, 40))
    ws.freeze_panes = "C2"
    stats["varian_baru"] = len(vb)

    # ── BOM_AKSESORIS (butuh isian) & BOM_OTOMATIS (sudah terisi otomatis — cukup periksa) ──
    ws = wb.create_sheet("BOM_AKSESORIS")
    _head(ws, BOM_FOKUS_COLS)
    ws_oto = wb.create_sheet("BOM_OTOMATIS")
    _head(ws_oto, BOM_FOKUS_COLS)
    in_file: set[str] = set()
    if parsed:
        rows = gap_sisa.bom_rows_from_file(data or b"")
        b_rows, b_styles, b_stats = await bom_fokus_rows(db, parsed, rows)
        manual, otomatis = split_otomatis(b_rows, b_styles)
        for target, pairs in ((ws, manual), (ws_oto, otomatis)):
            for r, st in pairs:
                target.append(r)
                i = target.max_row
                for col, fill in st["fills"]:
                    target.cell(row=i, column=BOM_FOKUS_COLS.index(col) + 1).fill = fill
        in_file = {clean(r[0]).upper() for r in rows if clean(r[0])} | {g["model_code"] for g in parsed.get("bom_groups") or []}
        stats.update({f"bom_{k}": v for k, v in b_stats.items()})
        stats["bom_manual_baris"] = len(manual)
        stats["bom_otomatis_baris"] = len(otomatis)
        stats["bom_otomatis_kelompok"] = len({st["group"] for _r, st in otomatis})
    # model yang belum punya aksesoris & tidak ada di berkas → 3 baris kosong
    models = await db.rahaza_models.find({"active": {"$ne": False}}, {"_id": 0, "id": 1, "code": 1, "name": 1}).sort("code", 1).to_list(5000)
    has_acc: set[str] = set()
    async for b in db.rahaza_boms.find({"active": {"$ne": False}, "is_active": True}, {"_id": 0, "model_id": 1, "materials": 1}):
        if any((ln.get("material_type") or "").lower() not in ("fabric", "") and not ln.get("is_cut_panel") for ln in b.get("materials") or []):
            has_acc.add(b["model_id"])
    vmap = await load_model_variants(db, [m["id"] for m in models])
    n_kosong = 0
    for m in models:
        if m["id"] in has_acc or m["code"] in in_file or not vmap.get(m["id"]):
            continue
        n_kosong += 1
        for k in range(3):
            ws.append([m["code"], m["name"], "", "", None, "", "", "", _variants_label(vmap[m["id"]]) if k == 0 else "",
                       "belum ada aksesoris di BOM — isi kode_material (lihat REF_AKSESORIS) & qty_per_pcs; tambah baris bila perlu" if k == 0 else ""])
            _fill(ws, ws.max_row, BOM_FOKUS_COLS, ["kode_material", "qty_per_pcs"], YELLOW)
    stats["bom_model_tanpa_aksesoris"] = n_kosong
    for sheet in (ws, ws_oto):
        for c in sheet[1]:
            c.alignment = WRAP
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
            row[BOM_FOKUS_COLS.index("yang_perlu_diisi")].alignment = WRAP
        _widths(sheet, (12, 22, 16, 40, 12, 8, 30, 26, 60, 80))
        sheet.freeze_panes = "C2"

    # ── MATERIAL: harga 0 + isi kemasan yang dibutuhkan BOM ──
    need_pack = gap_sisa.materials_needing_pack(parsed) if parsed else {}
    ws = wb.create_sheet("MATERIAL")
    _head(ws, MAT_FOKUS_COLS)
    mats = await db.rahaza_materials.find({"active": {"$ne": False}, "type": {"$ne": "fg"}, "code": {"$not": {"$regex": "^CUT-"}}},
                                          {"_id": 0}).sort("code", 1).to_list(20000)
    n_mat = 0
    for m in mats:
        price = float(m.get("unit_cost") or 0)
        no_price, pack = price <= 0, m.get("code") in need_pack
        if not no_price and not pack:
            continue
        base = m.get("unit") or ""
        todo = []
        if no_price:
            todo.append(f"isi harga_per_satuan_beli = harga 1 {base}")
        if pack:
            todo.append(f"isi isi_per_satuan_beli = berapa pcs dalam 1 {base} (dipakai BOM {need_pack[m['code']].split('dipakai BOM ')[-1]})")
        ws.append([m.get("code"), m.get("name"), m.get("type"), m.get("category_name") or m.get("category"), base,
                   base, None if pack else 1, None if no_price else price, price, m.get("min_stock") or None, "; ".join(todo)])
        i = ws.max_row
        _fill(ws, i, MAT_FOKUS_COLS, ["satuan_beli"], BLUE)
        _fill(ws, i, MAT_FOKUS_COLS, ["isi_per_satuan_beli"], YELLOW if pack else BLUE)
        _fill(ws, i, MAT_FOKUS_COLS, ["harga_per_satuan_beli"], YELLOW if no_price else BLUE)
        n_mat += 1
    _widths(ws, (14, 42, 10, 16, 12, 12, 18, 22, 22, 10, 70))
    ws.freeze_panes = "C2"
    stats["material"] = n_mat

    # ── REF_AKSESORIS (referensi) ──
    ws = wb.create_sheet("REF_AKSESORIS")
    _head(ws, ["kode_material", "nama", "tipe", "satuan_dasar", "harga_per_satuan_dasar"])
    for m in mats:
        if (m.get("type") or "").lower() != "fabric":
            ws.append([m.get("code"), m.get("name"), m.get("type"), m.get("unit"), float(m.get("unit_cost") or 0)])
    _widths(ws, (16, 44, 12, 12, 18))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), stats
